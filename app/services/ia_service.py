import os
import openai
from PyPDF2 import PdfReader
from docx import Document
import openpyxl  # Para .xlsx

from nltk.tokenize import sent_tokenize
from sumy.parsers.plaintext import PlaintextParser
from sumy.summarizers.lsa import LsaSummarizer

def ler_arquivo_texto(caminho, ext):
    if ext == '.pdf':
        reader = PdfReader(caminho)
        texto = ""
        for page in reader.pages:
            texto += page.extract_text() or ""
        return texto[:4000]
    elif ext in ['.doc', '.docx']:
        doc = Document(caminho)
        return ' '.join([p.text for p in doc.paragraphs])[:4000]
    elif ext in ['.txt']:
        with open(caminho, encoding='utf-8') as f:
            return f.read()[:4000]
    elif ext == '.xlsx':
        wb = openpyxl.load_workbook(caminho, data_only=True)
        texto = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        texto.append(str(cell))
        return ' '.join(texto)[:4000]
    return ""

USE_OPENAI = os.getenv("USE_OPENAI", "false").lower() == "true"

def gerar_resumo_local(texto, sentencas=3):
    sentences = sent_tokenize(texto, language="english")
    texto_tokenizado = "\n".join(sentences)
    parser = PlaintextParser.from_string(texto_tokenizado, lambda txt: txt.split('\n'))
    summarizer = LsaSummarizer()
    result = summarizer(parser.document, sentencas)
    return " ".join([str(s) for s in result]) or "Resumo local não disponível."

def gerar_resumo_ia(texto):
    if not texto:
        return ""
    if USE_OPENAI:
        try:
            client = openai.OpenAI()
            chat_completion = client.chat.completions.create(
                model="gpt-3.5-turbo",
                messages=[
                    {"role": "system", "content": "Resuma o texto enviado de forma clara e objetiva, em português."},
                    {"role": "user", "content": texto}
                ],
                max_tokens=150
            )
            return chat_completion.choices[0].message.content.strip()
        except Exception as e:
            print("Erro IA OpenAI:", e)
            return gerar_resumo_local(texto)
    else:
        return gerar_resumo_local(texto)
